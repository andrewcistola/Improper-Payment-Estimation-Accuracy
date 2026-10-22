"""Build qa/QA_Checklist.xlsx, the human reviewer's checklist.

Run from the repository's top folder:  python qa/scripts/build_checklist.py
(then recalculate with LibreOffice so the Summary formulas carry values).
Edit the rows below, not the .xlsx, so the checklist stays traceable.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "qa" / "QA_Checklist.xlsx"

PY = ".venv/bin/python"   # Windows: .venv\Scripts\python
STATUSES = ["Not started", "Pass", "Fail", "N/A"]

# --------------------------------------------------------------------------
# Checklist rows: (section, task, how to check, where to look, expected result)
# --------------------------------------------------------------------------
SECTIONS = {
    "A": "Setup and independence",
    "B": "Re-run with date-stamped logs",
    "C": "Code review: key assumptions",
    "D": "Results: represented correctly",
    "E": "Tests and independent checks",
    "F": "Data security and disclosure",
    "G": "Documentation and reproducibility",
    "H": "Close-out and sign-off",
}

ROWS = [
    # ---------------- A. Setup and independence ----------------
    ("A", "Record who you are and your relationship to the study",
     "Fill in the Reviewer sheet. State whether you contributed to the study's code, data, or writing.",
     "Reviewer sheet", "Reviewer sheet complete; any relationship to the authors disclosed"),
    ("A", "Use a fresh copy of the project",
     "Download the project from GitHub (ZIP or git clone) into a new, empty folder. Do not use the author's computer or an existing copy.",
     "start-here/README.md, step 3", "A new project folder on your own computer"),
    ("A", "Record the project version you are reviewing",
     "Git: run  git rev-parse HEAD  in the project folder. ZIP: record the download date and time. The re-run also records it in environment.txt ('Project version').",
     "Reviewer sheet; start-here/logs/run_*/environment.txt", "Commit ID (or ZIP download date) recorded"),
    ("A", "Install the software exactly as instructed",
     "Follow start-here/README.md steps 1 to 5 without changes. Note any step you had to change or that was unclear.",
     "start-here/README.md, steps 1-5", "Setup ends with 'Successfully installed'; deviations noted"),
    ("A", "Confirm the pinned DuckDB version",
     f"Run  {PY} -c \"import duckdb; print(duckdb.__version__)\"",
     "pyproject.toml; requirements.txt", "1.5.6 (random draws depend on this version)"),

    # ---------------- B. Re-run with date-stamped logs ----------------
    ("B", "Run the quick trial",
     f"Run  {PY} start-here/run_all.py --quick  and record the log folder name.",
     "start-here/logs/run_<date>_<time>_quick/", "Ends with 'Your setup works.'"),
    ("B", "Start the full replication and record the start date and time",
     f"Run  {PY} start-here/run_all.py . Record the start time on the Run log sheet.",
     "start-here/logs/run_<date>_<time>/run_log.txt (first line)", "Run started; log folder created"),
    ("B", "Stage 1 completed: populations built and calibration checks passed",
     "In run_log.txt, find 'Stage 1 finished'. Open each errors/output*/.../checks.json and confirm every check passes.",
     "run_log.txt; */errors/output*/*/checks.json", "Stage 1 finished; all calibration checks pass"),
    ("B", "Stage 2 completed: reported estimates (parametric-rate bootstrap)",
     "Find 'Stage 2 finished' in run_log.txt. Record the time on the Run log sheet.",
     "run_log.txt; */estimation/results/", "Stage 2 finished, no FAILED lines"),
    ("B", "Stage 3 completed: fixed-rate comparison",
     "Find 'Stage 3 finished' in run_log.txt.", "run_log.txt; */estimation/results_fixed/",
     "Stage 3 finished, no FAILED lines"),
    ("B", "Stage 4 completed: Part D pre-IRA reference",
     "Find 'Stage 4 finished' in run_log.txt.", "run_log.txt; part-d-ipm/pre-ira-reference.md",
     "Stage 4 finished, no FAILED lines"),
    ("B", "Stages 5 and 6 completed: report and manuscript tables",
     "Find 'Stage 5 finished' and 'Stage 6 finished' in run_log.txt.",
     "run_log.txt; shared/results/report/; manuscripts/generated/", "Both stages finished"),
    ("B", "Stage 7 completed: automated tests pass",
     "Open the stage7 log. The last pytest line reports the number passed.",
     "start-here/logs/run_*/stage7_01_pytest.log", "All tests passed, 0 failed"),
    ("B", "Stage 8: the result check reproduces the published results",
     "Open the stage8 log (or verification/verify_check.txt after collecting).",
     "start-here/logs/run_*/stage8_01_start_here_verify.log",
     "REPRODUCED EXACTLY (or REPRODUCED, with the reason explained in Notes)"),
    ("B", "If the run was interrupted, every stage still completed once",
     "If you resumed with --from, confirm that each stage 1-8 has a 'Stage N finished' line across the run folders.",
     "start-here/logs/run_*/run_log.txt", "Every stage finished at least once, in order; or N/A"),
    ("B", "The logs contain no errors",
     "Search every .log file for 'Traceback', 'Error', 'FAILED' and 'exit code: 1'. (A RuntimeWarning is not an error.)",
     "start-here/logs/run_*/*.log", "None found, or each explained in Notes"),
    ("B", "Collect the re-run into the QA holding space",
     f"Run  {PY} qa/collect_rerun.py --reviewer \"Your Name\" --initials XX",
     "qa/reruns/<date>_<initials>/", "Folder created; RERUN_SUMMARY.md lists your runs and the result check"),
    ("B", "Review the environment record",
     "Open environment.txt in each run folder. Confirm the date, operating system, Python and package versions are yours.",
     "qa/reruns/<date>_<initials>/logs/run_*/environment.txt", "Matches your computer; copied to the Reviewer sheet"),
    ("B", "Review the file inventory",
     "Open file_inventory.csv in Excel. Filter matches_published for 'NO'.",
     "qa/reruns/<date>_<initials>/verification/file_inventory.csv", "No rows with 'NO', or each explained in Notes"),

    # ---------------- C. Code review: key assumptions ----------------
    # Common to all programs
    ("C", "Review is perfect: the audit finding equals the true error",
     "Confirm that the sampled units' y is read directly from the population's true y, with no review error or appeals.",
     "shared/sql/sample_srs.sql; shared/sql/sample_systematic.sql; */estimation/sql/*01_frame.sql",
     "Sample y = population y; no reviewer error modeled (manuscript Appendix B.6, item 1)"),
    ("C", "One audit per audit year; every method uses only that audit",
     "Confirm that expansion, ratio and regression read the same sample table, and that the bootstrap simulation uses only that year's audit findings (Part D: its own 5,000-PDE audit).",
     "*/estimation/config.yaml (step order); shared/ipsim/simulation.py",
     "No method sees the population's true errors; Part D's two audits are as documented"),
    ("C", "Errors are independent of selection within a stratum",
     "Confirm that errors are fabricated before sampling and that sample selection uses only stratum and unit id (not y).",
     "*/errors/sql/*apply_errors*.sql; shared/sql/sample_srs.sql; shared/sql/lib/00_rng.sql",
     "Selection keys never use y"),
    ("C", "Neyman allocation uses the population's stratum standard deviations",
     "Confirm n_h is proportional to N_h x S_h, with S_h computed from the population (standing in for prior-year data), and sums to the stated audit size.",
     "eipm/estimation/sql/01_frame.sql; perm/estimation/sql/01_frame.sql, cert_01_frame.sql; part-d-ipm/estimation/sql/40_simulation_inputs.sql; allocation.json",
     "Neyman formula; sizes 2,000 (EIPM), 520 per state (PERM), 37,651 (CERT), 5,000 PDEs (Part D)"),
    ("C", "Netting: errors netted within the payment unit before measures",
     "Confirm over = sum max(y,0), under = sum max(-y,0), gross = over + under, net = over - under, after netting per enrollee/beneficiary/claim.",
     "*/estimation/sql/*01_frame.sql; */errors/sql/*truth*.sql",
     "Matches manuscript Section 3.3"),
    ("C", "Unknown payments (e.g., missing documentation) count as improper",
     "Confirm missing-documentation errors set the whole payment (or every HCC) to improper.",
     "part-c-ipm/errors/sql/04_apply_errors.sql; cert/errors/sql/02_apply_errors.sql",
     "Follows CMS practice (Appendix B.6, item 6)"),
    ("C", "Calibration is one closed-form multiplier, with no iterative tuning",
     "Confirm each calibrated rate = base probability x (target dollars / expected dollars), and that the multiplier is logged.",
     "*/errors/sql/*apply_errors*.sql; part-d-ipm/errors/sql/03_init_calibration.sql; truth.json (by.calibration)",
     "One multiplier per target; recorded in truth.json"),
    ("C", "Random draws are deterministic functions of (seed, stream, key)",
     "Confirm SQL draws hash a single string per stream and simulation draws use PCG64 seeded by SeedSequence([scenario seed, audit year]).",
     "shared/sql/lib/00_rng.sql; shared/ipsim/engine.py (child_seed); shared/ipsim/simulation.py",
     "No unseeded randomness anywhere"),
    ("C", "Scale factors convert synthetic units to national dollars",
     "Confirm w_pop: Part C 100, Part D 500, EIPM 100, PERM 200, CERT 200; national dollars = rate x total payments x w_pop.",
     "*/errors/config.yaml (w_pop); shared/sql/report.sql",
     "Values as stated; rates are scale-free"),
    ("C", "Analytic intervals: normal (1.96) with finite population correction",
     "Confirm expansion and ratio variances use the stratified SRS formula with (1 - n_h/N_h), and intervals are estimate +/- 1.96 SE.",
     "shared/sql/lib/02_estimators.sql (ipm_expansion, ipm_ratio)", "Matches manuscript Section 3.4"),
    ("C", "Regression: models fit within strata; residual t interval",
     "Confirm design-weighted least squares within each model stratum, T = sum of predictions + weighted residuals, variance from residuals, t with df = n - strata - terms.",
     "shared/ipsim/regression.py; */estimation/sql/*30_regression.sql",
     "Matches manuscript Section 3.4"),
    ("C", "Regression fallback: intercept-only with too few errors",
     "Confirm a stratum uses an intercept-only model when it has fewer than 10 audited errors or fewer than (terms + 5) units.",
     "shared/ipsim/regression.py; */estimation/config.yaml (reg_min_errors: 10)", "Threshold 10 errors; terms + 5 units"),
    ("C", "Regression predictors by program",
     "Confirm reg_predictors: Part C x, demo_score, months, n_hcc; Part D x, gdc_sub, n_pde, catastrophic (post-strata LTC/LIS/other); EIPM x, age, hh_size, fpl_pct, months; PERM x, month, bene_n_pay, bene_paid; CERT x.",
     "*/estimation/config.yaml", "Predictors as listed in manuscript Section 3.4"),
    ("C", "Bootstrap simulation (fixed rates): p = f/n, q = g/f per donor kind",
     "Confirm each non-audited unit is in error with p_k, partial with q_k, partial size drawn uniformly from the audited partial sizes.",
     "shared/ipsim/simulation.py; */estimation/*simulation.py; */estimation/sql/*40_simulation_inputs.sql",
     "Matches manuscript Section 3.5"),
    ("C", "Bootstrap simulation (parametric rates): binomial redraw each iteration",
     "Confirm p~ = Binomial(n, p)/n and q~ = Binomial(f, q)/f before every iteration, and that zero-failure rates stay at zero.",
     "shared/ipsim/simulation.py (sim_rates: parametric)", "Matches manuscript Section 3.6"),
    ("C", "Audited units keep their own findings",
     "Confirm the simulation assigns errors only to non-audited units and keeps each audited unit's actual finding.",
     "*/estimation/*simulation.py", "Audited units unchanged across iterations"),
    ("C", "Simulation estimate and interval: mean and 2.5th-97.5th percentiles of 1,000 iterations",
     "Confirm sim_iterations: 1000 and the percentile interval.", "shared/ipsim/simulation.py; */estimation/config.yaml",
     "1,000 iterations; mean; 2.5th and 97.5th percentiles"),
    ("C", "Scoring: coverage, 25 audit years, medians",
     "Confirm an interval covers if lower <= truth <= upper, undefined intervals count as not covering, and the report uses audit years 1-25 with medians.",
     "shared/sql/lib/02_estimators.sql (ipm_score); shared/sql/report.sql; shared/report.yaml (audit_years: 25)",
     "As stated in the README and manuscript Section 3.7"),
    # Part C
    ("C", "Part C: population size, payment total and HCC model",
     "Confirm about 330,000 enrollees (x100), 550 contracts, payment scaled to $388.72B, HCC hierarchies (severe/mild within 10 illustrative families).",
     "part-c-ipm/errors/config.yaml; part-c-ipm/errors/sql/01-03*.sql", "Matches Appendix B.1"),
    ("C", "Part C: error mechanisms and calibration targets",
     "Confirm unsupported HCCs (deleted or moved down), missing documentation, and supported-but-not-coded HCCs (underpayments); targets 5.29% unsupported, 0.22% missing documentation, 0.57% underpayments.",
     "part-c-ipm/errors/config.yaml (target_*); part-c-ipm/errors/sql/04_apply_errors.sql",
     "FY 2025 Part C: 6.09% gross, $21.43B over, $2.23B under"),
    ("C", "Part C: audit design (990 enrollees, 330 per risk-score tercile)",
     "Confirm the estimation allocation. NOTE: the design block in part-c-ipm/errors/config.yaml lists 1,500 per stratum; confirm it is description only and not used by the estimation.",
     "part-c-ipm/estimation/config.yaml (allocation); part-c-ipm/errors/config.yaml (design)",
     "330 x 3 used; record the design-block discrepancy as an issue (Minor)"),
    ("C", "Part C: expansion in risk-score units; simulation donors are HCCs",
     "Confirm expansion = mean risk-score error x tercile size x national mean county rate x months; simulation uses 20 HCC failure rates and 10 family addition rates, never extrapolating to unaudited HCCs.",
     "part-c-ipm/estimation/sql/10_expansion.sql; part-c-ipm/estimation/simulation.py; sql/40_simulation_inputs.sql",
     "Matches Appendix B.1"),
    # Part D
    ("C", "Part D: post-IRA CY 2025 benefit parameters",
     "Confirm $590 deductible, $2,000 out-of-pocket threshold, 25% initial coinsurance, manufacturer 10%/20% (brand), reinsurance 20% brand/40% generic, LIS copays $4.90/$12.15.",
     "part-d-ipm/errors/config.yaml; part-d-ipm/errors/sql/05_benefit.sql", "Matches Appendix B.2 and CMS CY 2025 parameters"),
    ("C", "Part D: error rates and sizes (not calibrated)",
     "Confirm 10% of LTC PDEs and 2.5% of other PDEs in error; 90% full (GDC to 0), 10% partial uniform on [-100%, 100%). Truth 1.38% (published FY 2025 rate 4.00% is intentionally not targeted).",
     "part-d-ipm/errors/config.yaml; part-d-ipm/errors/sql/02_candidate_errors.sql, 04_apply_errors.sql",
     "Matches Appendix B.2"),
    ("C", "Part D: numpy payment function reproduces the SQL benefit to the cent",
     "Confirm the test exists and passed in stage 7.", "shared/tests/test_simulation.py; part-d-ipm/estimation/simulation.py",
     "Test present and passing"),
    ("C", "Part D: two audits and the expansion delta-method interval",
     "Confirm the 990-beneficiary SRS (ratio, regression) and the 5,000-PDE Neyman audit in 3 strata (expansion, simulation); expansion interval moves each stratum rate +/-1 SE and reruns the benefit.",
     "part-d-ipm/estimation/config.yaml; sql/40_simulation_inputs.sql; simulation.py (run_expansion)",
     "Matches Appendix B.2"),
    ("C", "Part D pre-IRA: CY 2023 benefit parameters",
     "Confirm $505 deductible, $4,660 initial coverage limit, $7,400 TrOOP threshold, 70% coverage-gap discount, 80% catastrophic reinsurance, LIS copays $4.15/$10.35.",
     "part-d-ipm/errors/pre_ira.yaml; part-d-ipm/errors/sql/05_benefit_pre_ira.sql", "Matches part-d-ipm/pre-ira-reference.md"),
    # PERM
    ("C", "PERM: states, error rates and federal share",
     "Confirm 9 states with eligibility rates on a 2.00%-4.00% grid; FFS claims errors 6.5% (70% all-or-nothing, 30% partial), MC 0.3%; rates at federal share (sum y x FMAP / sum x x FMAP).",
     "perm/errors/config.yaml; perm/errors/sql/02*.sql, perm_03*.sql", "Matches Appendix B.3; pooled about 6.12%"),
    ("C", "PERM: gross combines eligibility and claims errors without double counting",
     "Confirm gross = whole payment if the beneficiary is ineligible, otherwise the claims error.",
     "perm/errors/sql/04_truth.sql", "Claims + eligibility - overlap"),
    ("C", "PERM: audit design and national pooling",
     "Confirm 520 payments per state, 10 strata (payment type x 5 quintiles), Neyman, systematic selection after sorting by amount; national pooling adds dollars and variances (analytic) and sums iteration totals (simulation).",
     "perm/estimation/config.yaml; shared/sql/sample_systematic.sql; perm/estimation/sql/95_national.sql",
     "Matches Appendix B.3"),
    # EIPM
    ("C", "EIPM: overpayments only; technical and monetary errors",
     "Confirm no underpayments; technical errors (whole APTC) are 48% of improper dollars; monetary errors raise supported income by a lognormal factor and recompute APTC; target 0.89%.",
     "eipm/errors/config.yaml; eipm/errors/sql/02_apply_errors.sql", "Matches Appendix B.4"),
    ("C", "EIPM: APTC schedule and simulation income clipping",
     "Confirm the enhanced applicable-percentage schedule (0% to 8.5%) and that simulated incomes are clipped to 100-600% FPL.",
     "eipm/errors/config.yaml (applicable_schedule); eipm/estimation/simulation.py", "As stated"),
    ("C", "EIPM: audit design",
     "Confirm stratified SRS of 2,000 applications, Neyman over 5 APTC quintiles (illustrative strata).",
     "eipm/estimation/config.yaml; eipm/estimation/sql/01_frame.sql", "Matches Appendix B.4"),
    # CERT
    ("C", "CERT: claim types and payments",
     "Confirm 4 claim types (IPPS 40,000; Part A other 900,000; Part B 4,000,000; DMEPOS 400,000 synthetic claims) scaled to $118.5B, $197.9B, $114.1B, $9.4B ($439.9B).",
     "cert/errors/config.yaml; cert/errors/sql/01_population.sql", "Matches Appendix B.5"),
    ("C", "CERT: error mix and calibration",
     "Confirm whole-claim overpayments 92% of overpayment dollars, partial 8% (5-60% of payment), underpayments 5-50%; per-type closed-form rates from FY 2025 dollars.",
     "cert/errors/config.yaml; cert/errors/sql/02_apply_errors.sql",
     "Truth 6.50% gross (published 6.55%); $0.92B under"),
    ("C", "CERT: audit design",
     "Confirm 37,651 claims (IPPS 8,750; Part A other 8,651; Part B 11,600; DMEPOS 8,650), Neyman across 5 payment quintiles within type; no MAC jurisdictions; no A-to-B rebilling.",
     "cert/estimation/config.yaml; sql/cert_01_frame.sql, cert_02_sample.sql", "Matches Appendix B.5"),
    # Reporting
    ("C", "Audit-size projection formula",
     "Confirm the share of the current audit = (parametric half-width / analytic half-width) squared.",
     "manuscripts/scripts/build_tables.py", "Matches manuscript Sections 3.7 and 5.2"),

    # ---------------- D. Results represented correctly ----------------
    ("D", "Main README: gross improper payments (dollars) table",
     "For each program and method, compare True / Estimated / Upper / Lower with by_scenario.csv (domain ending in 'gross'; PERM and CERT 'national' rows): true_improper_usd, median_estimate_usd, median_upper_usd, median_lower_usd.",
     "README.md (Results); shared/results/report/by_scenario.csv", "Every value matches after rounding to $0.01B"),
    ("D", "Main README: gross improper payment rates table",
     "Same comparison using true_rate, median_estimate, median_upper, median_lower.",
     "README.md (Results); shared/results/report/by_scenario.csv", "Every value matches after rounding to 0.01%"),
    ("D", "Main README: Findings and recommendations numbers",
     "Check 121/125 (96.8%), 111/125, 108/125, 30/125, 11-21% narrower, 63%, 1.5-6.6% and 16.6% MAE, 19-20 of 25, 15 of 25, and 20-38% smaller audits against the tables.",
     "README.md (Findings); manuscripts/generated/tables.md; results_long.csv", "Every number traceable to a table"),
    ("D", "Program READMEs: results tables",
     "For each program README, compare the Dollars, Rates and coverage tables with by_scenario.csv.",
     "part-c-ipm, part-d-ipm, eipm, perm, cert README.md; by_scenario.csv", "All values match"),
    ("D", "Manuscript tables match the generated tables",
     "Compare Tables 2, 3, 4, 5 and C1 in manuscript.pdf with tables.md.",
     "manuscripts/manuscript.pdf; manuscripts/generated/tables.md", "Identical"),
    ("D", "Manuscript text numbers match the tables",
     "Check every number in the abstract, Results (Section 4) and Discussion (Section 5) against Tables 2-5.",
     "manuscripts/sections/*.md; manuscripts/generated/tables.md", "Every number traceable; none contradicts a table"),
    ("D", "Manuscript Table 1 (populations) matches the configs and truth files",
     "Compare unit counts, audit sizes and true gross rates with the configs and truth.json.",
     "manuscripts/sections/01_background_methods.md; */errors/output*/*/truth.json", "All match"),
    ("D", "Figures 1 and 2 show the values in results_long.csv",
     "Spot-check at least 3 bars in each figure against results_long.csv.",
     "manuscripts/generated/fig1_width.png, fig2_coverage.png; results_long.csv", "Bars match the table values"),
    ("D", "One-pager and conference abstract numbers",
     "Check every number against the manuscript tables; confirm the abstract is 500 words or fewer.",
     "manuscripts/one-pager.md; manuscripts/conference-abstract.md", "All match"),
    ("D", "Calibration targets match public sources",
     "Check the FY 2025 public figures used as targets (Part C 6.09%, $21.43B/$2.23B; Medicaid 6.12%; EIPM 0.89% with 0.43 technical; CERT 6.55% and claim-type figures; Part D 4.00%) against the cited CMS sources.",
     "manuscript Section 3.2 and Appendix B; literature review Section 3 links", "Each target matches its public source"),
    ("D", "True values in the tables equal the truth files",
     "Confirm true_rate in by_scenario.csv equals theta in each estimation truth.json and the populations' truth.json.",
     "by_scenario.csv; */estimation/results/*/truth.json; */errors/output*/*/truth.json", "Identical"),
    ("D", "Hand-check one median and one coverage count",
     "Open by_audit_year.csv (or .parquet). Filter Part C, national|gross, simulation. Compute the median of the 25 estimates and count truth_in_ci = 'inside'.",
     "shared/results/report/by_audit_year.csv", "Median equals median_estimate; count equals n_truth_inside (24)"),
    ("D", "Part D pre-IRA reference is stated as reference only",
     "Confirm pre-ira-reference.md says it is not part of the reported results and that README tables use post-IRA results.",
     "part-d-ipm/pre-ira-reference.md; README.md", "Clearly labeled as a robustness reference"),
    ("D", "Conclusions are supported by the results",
     "Read the Conclusions and Recommendations. Confirm each claim (validity, accuracy, precision, audit-size reduction 'pending confirmation') is supported and its limits are stated.",
     "README.md; manuscript Sections 5-7", "No claim goes beyond the evidence"),
    ("D", "Limitations are complete",
     "Confirm the limitations cover 25 audit years (Monte Carlo SE about 4.4 points), synthetic populations, perfect review, Part D design differences, fixed error sizes, and one-sided recovery bounds untested.",
     "manuscript Section 5.4", "All present"),

    # ---------------- E. Tests and independent checks ----------------
    ("E", "Automated tests pass on your computer",
     f"Run  {PY} -m pytest -q shared/tests  (also run in stage 7).", "shared/tests/", "All passed"),
    ("E", "Results do not depend on the number of processor threads",
     f"Run  {PY} -m ipsim eipm/estimation/config.yaml --out qa/tmp_threads1 --threads 1 . Compare qa/tmp_threads1/base/manifest.json with eipm/estimation/results/base/manifest.json (estimates.parquet sha256). Then delete qa/tmp_threads1.",
     "eipm/estimation/results/base/manifest.json", "Identical sha256 fingerprints"),
    ("E", "Census check: auditing every unit returns the true rate",
     "Confirm the test exists and passed.", "shared/tests/test_simulation.py", "Test present and passing"),
    ("E", "Estimators match an independent implementation",
     "Confirm the SQL estimators are tested against a separate numpy implementation.", "shared/tests/test_estimators.py",
     "Test present and passing"),
    ("E", "Spot-check one analytic estimate by hand",
     "In by_audit_year.csv, pick EIPM national|gross, ratio, audit year 1. Confirm lower = estimate - 1.96 x se is consistent with the reported interval (se in estimates.parquet).",
     "eipm/estimation/results/base/estimates.parquet; by_audit_year.csv", "Consistent to rounding"),

    # ---------------- F. Data security and disclosure ----------------
    ("F", "No protected health information or real unit-level data",
     "Confirm every population is generated from stated distributions (errors/sql) and no data files are read from outside the project.",
     "*/errors/sql/*.sql; .gitignore (data/raw/ excluded)", "Only synthetic data"),
    ("F", "No non-public material",
     "Confirm there are no contractor methodology plans, internal error data, or non-public CMS documents; non-public design values are labeled illustrative.",
     "configs (comments 'illustrative', 'NOT public'); literature-review references", "Public sources only; illustrative values labeled"),
    ("F", "Sources for public figures are cited",
     "Confirm each public figure used for calibration is cited with a source.", "manuscript References; literature review References",
     "Every figure has a source"),
    ("F", "Author disclosures and approvals",
     "Confirm the funding, conflicts of interest and disclaimer section is completed, and any employer or agency clearance needed before publication has been obtained.",
     "manuscripts/sections/03_future_conclusion_refs.md", "Completed before submission (or noted as open)"),
    ("F", "License and repository contents",
     "Confirm a LICENSE file is present and the repository contains no credentials, personal data or large unexplained files.",
     "LICENSE; repository root", "As expected"),

    # ---------------- G. Documentation and reproducibility ----------------
    ("G", "Start Here instructions worked as written",
     "Record any step that failed, was unclear, or needed a change on your operating system.",
     "start-here/README.md", "No changes needed, or each recorded on the Issues log"),
    ("G", "Run times are close to the published estimates",
     "Compare your stage times (Run log sheet) with start-here/README.md step 7.", "start-here/README.md; Run log sheet",
     "Recorded; large differences explained"),
    ("G", "Every result traces to code and data in the repository",
     "Pick 3 numbers from the manuscript and follow each back to the file and code that produced it.",
     "manuscripts; shared/results; */estimation", "Each traced without gaps"),
    ("G", "Data dictionary describes the output files",
     "Spot-check 5 columns in estimates.parquet and by_scenario.csv against the data dictionary.",
     "shared/data_dictionary.md", "Definitions match"),

    # ---------------- H. Close-out and sign-off ----------------
    ("H", "Every Fail is on the Issues log",
     "Check that each row marked Fail has an issue ID in the Issue ID column.", "Checklist; Issues log", "Yes"),
    ("H", "Issues are resolved or accepted",
     "Confirm no Critical or Major issue is Open. Minor issues may remain open if noted on the sign-off form.",
     "Issues log; Summary sheet", "No open Critical or Major issues"),
    ("H", "Completed checklist saved in the re-run folder",
     "Save this workbook as QA_Checklist_<date>_<initials>.xlsx in qa/reruns/<date>_<initials>/review/.",
     "qa/reruns/<date>_<initials>/review/", "Saved"),
    ("H", "Sign-off form completed and saved",
     "Complete qa/QA_Signoff_Template.docx, sign it, and save it (and a PDF) in the review/ folder.",
     "qa/reruns/<date>_<initials>/review/", "Signed and saved"),
    ("H", "Re-run folder submitted",
     "Submit the qa/reruns/<date>_<initials>/ folder to the project (pull request or ZIP to the author).",
     "qa/README.md, step 9", "Submitted; date recorded"),
]

# --------------------------------------------------------------------------
ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def font(**kw):
    return Font(name=ARIAL, size=kw.pop("size", 10), **kw)


def header(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14, color="1F3864")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, color="595959")


wb = Workbook()

# ---------------- How to use ----------------
ws = wb.active
ws.title = "How to use"
title(ws, "QA Checklist: Comparative Accuracy of Improper Payment Estimation Methods",
      "Human reviewer's checklist. Read qa/README.md first.")
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 100
how = [
    ("Purpose", "This workbook records everything a human reviewer checks: an independent re-run with date-stamped logs, a code review of every key assumption, and a check that the results are represented correctly. It is the evidence behind the signed sign-off form."),
    ("Order of work", "1. Reviewer sheet  2. Re-run (Checklist section B) and Run log sheet  3. Code review (section C)  4. Results (section D)  5. Tests, data security, documentation (E-G)  6. Issues log  7. Close-out (section H) and sign-off form."),
    ("Cells you fill in", "Only the yellow cells. On the Checklist: Status, Date checked, Initials, Evidence, Issue ID and Notes. Grey and blue cells are fixed text or formulas; do not edit them."),
    ("Status values", "Not started: not yet checked.  Pass: checked and matches the expected result.  Fail: checked and does not match; log it on the Issues log and enter its ID.  N/A: does not apply (explain why in Notes)."),
    ("Dates", "Enter dates as YYYY-MM-DD (for example 2026-10-21). Date-times on the Run log as YYYY-MM-DD HH:MM."),
    ("Evidence", "Where the proof is: a file path in your re-run folder (for example logs/run_2026-10-21_093000/stage7_01_pytest.log), a screenshot file name in review/, or a short description."),
    ("Commands", "Commands are written for macOS and Linux (.venv/bin/python). On Windows, replace .venv/bin/python with .venv\\Scripts\\python and forward slashes in file paths with backslashes."),
    ("Wildcards", "In 'Where to look', * means 'any folder or file name here'; for example */errors/sql/*.sql means the errors/sql folder of every program."),
    ("Example of a completed row", "Status: Pass | Date checked: 2026-10-21 | Initials: JD | Evidence: logs/run_2026-10-21_093000/stage7_01_pytest.log | Issue ID: (blank) | Notes: 38 passed in 89 s."),
    ("When you finish", "The Summary sheet shows progress and whether the review is ready to sign. Save this file as QA_Checklist_<date>_<initials>.xlsx in qa/reruns/<date>_<initials>/review/."),
]
for i, (k, v) in enumerate(how, 4):
    ws.cell(row=i, column=1, value=k).font = font(bold=True)
    c = ws.cell(row=i, column=2, value=v)
    c.font = font()
    c.alignment = WRAP
    ws.cell(row=i, column=1).alignment = WRAP
leg = len(how) + 5
ws.cell(row=leg, column=1, value="Legend").font = font(bold=True)
ws.cell(row=leg + 1, column=1, value="Yellow").fill = INPUT_FILL
ws.cell(row=leg + 1, column=2, value="You fill this in").font = font()
ws.cell(row=leg + 2, column=1, value="Light blue").fill = SEC_FILL
ws.cell(row=leg + 2, column=2, value="Section heading or fixed text").font = font()
ws.cell(row=leg + 3, column=1, value="Green / red status").font = font()
ws.cell(row=leg + 3, column=2, value="Pass turns green; Fail turns red automatically").font = font()

# ---------------- Reviewer ----------------
rv = wb.create_sheet("Reviewer")
title(rv, "Reviewer and environment", "Fill in every yellow cell. Copy computer and software details from environment.txt in your run folder.")
rv.column_dimensions["A"].width = 44
rv.column_dimensions["B"].width = 60
rv.column_dimensions["C"].width = 60
fields = [
    ("REVIEWER", None, None),
    ("Full name", "", ""),
    ("Initials (used in the re-run folder name)", "", ""),
    ("Title / role", "", ""),
    ("Organization", "", ""),
    ("Email", "", ""),
    ("Relationship to the study or its authors", "", "For example: none; colleague; co-author. Disclose any role in writing the code or text."),
    ("Independent of the study's development? (Yes/No)", "", ""),
    ("REVIEW", None, None),
    ("Review start date (YYYY-MM-DD)", "", ""),
    ("Review end date (YYYY-MM-DD)", "", ""),
    ("Project version reviewed (Git commit ID, or ZIP download date)", "", "From environment.txt, 'Project version'"),
    ("How the project was obtained (Git clone / ZIP download)", "", ""),
    ("Re-run folder (qa/reruns/...)", "", "For example qa/reruns/2026-10-21_JD"),
    ("COMPUTER AND SOFTWARE (from environment.txt)", None, None),
    ("Computer (make/model, or 'personal laptop', 'work desktop')", "", ""),
    ("Operating system and version", "", ""),
    ("Processor cores", "", ""),
    ("Memory (GB)", "", ""),
    ("Python version", "", ""),
    ("DuckDB version (must be 1.5.6)", "", ""),
    ("numpy / numba / pyarrow / scipy versions", "", ""),
    ("TOOLS USED FOR THE REVIEW", None, None),
    ("Spreadsheet software (e.g., Microsoft Excel 365, version)", "", ""),
    ("Text editor or code viewer (e.g., VS Code, GitHub website)", "", ""),
    ("Terminal (e.g., PowerShell, macOS Terminal)", "", ""),
    ("PDF / Word viewer", "", ""),
    ("Other tools (e.g., Git, R, calculator)", "", ""),
]
rv["A3"], rv["B3"], rv["C3"] = "Item", "Your entry", "Guidance"
for c in ("A3", "B3", "C3"):
    rv[c].font = font(bold=True, color="FFFFFF")
    rv[c].fill = HEAD_FILL
r = 4
for lab, val, guide in fields:
    a = rv.cell(row=r, column=1, value=lab)
    if val is None:
        a.font = font(bold=True)
        for j in (1, 2, 3):
            rv.cell(row=r, column=j).fill = SEC_FILL
    else:
        a.font = font()
        b = rv.cell(row=r, column=2)
        b.fill = INPUT_FILL
        b.border = BOX
        b.font = font()
        g = rv.cell(row=r, column=3, value=guide)
        g.font = font(italic=True, color="595959")
        g.alignment = WRAP
    a.alignment = WRAP
    r += 1
rv.freeze_panes = "A4"

# ---------------- Checklist ----------------
ck = wb.create_sheet("Checklist")
title(ck, "Checklist", "One row per task. Fill in the yellow columns. Fail rows need an Issue ID from the Issues log.")
cols = ["ID", "Section", "Task", "How to check", "Where to look", "Expected result",
        "Status", "Date checked", "Initials", "Evidence (file, screenshot, or description)", "Issue ID", "Notes"]
header(ck, 4, cols, [7, 22, 38, 60, 44, 36, 13, 13, 9, 36, 10, 36])
row = 5
counters = {}
first_data, last_data = None, None
current = None
for sec, task, how_, where, expect in ROWS:
    if sec != current:
        current = sec
        c = ck.cell(row=row, column=1, value=f"{sec}. {SECTIONS[sec]}")
        c.font = font(bold=True, color="1F3864")
        for j in range(1, len(cols) + 1):
            ck.cell(row=row, column=j).fill = SEC_FILL
        row += 1
    counters[sec] = counters.get(sec, 0) + 1
    vals = [f"{sec}{counters[sec]}", SECTIONS[sec], task, how_, where, expect, "Not started", None, None, None, None, None]
    for j, v in enumerate(vals, 1):
        c = ck.cell(row=row, column=j, value=v)
        c.font = font()
        c.alignment = WRAP
        c.border = BOX
        if j >= 7:
            c.fill = INPUT_FILL
    ck.cell(row=row, column=8).number_format = "yyyy-mm-dd"
    first_data = first_data or row
    last_data = row
    row += 1
ck.freeze_panes = "D5"
ck.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{last_data}"
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False,
                    showErrorMessage=True, errorTitle="Status", error="Choose Not started, Pass, Fail or N/A.")
ck.add_data_validation(dv)
dv.add(f"G{first_data}:G{last_data}")
rng = f"G{first_data}:G{last_data}"
ck.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ck.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor="FFC7CE")))
ck.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N/A"'], fill=PatternFill("solid", fgColor="EDEDED")))
ck.conditional_formatting.add(f"K{first_data}:K{last_data}",
                              FormulaRule(formula=[f'AND($G{first_data}="Fail",$K{first_data}="")'],
                                          fill=PatternFill("solid", fgColor="FFC7CE")))
ck["K4"].comment = Comment("Required for every row marked Fail: the matching ID from the Issues log (e.g., I-01).", "QA")

# ---------------- Run log ----------------
rl = wb.create_sheet("Run log")
title(rl, "Run log", "One row per run folder in start-here/logs/ (copy from RERUN_SUMMARY.md and run_log.txt). Duration is calculated.")
rcols = ["Run folder name", "Run type", "Stages run", "Started (YYYY-MM-DD HH:MM)", "Finished (YYYY-MM-DD HH:MM)",
         "Duration (h:mm)", "Outcome", "Result check", "Notes"]
header(rl, 4, rcols, [36, 18, 12, 22, 22, 14, 18, 24, 44])
ex = ["run_2026-10-21_093000_quick (EXAMPLE: replace)", "Quick trial", "1-2", "2026-10-21 09:30", "2026-10-21 09:33",
      None, "Finished", "N/A (trial)", "Example row showing the format; overwrite it with your first run."]
for r in range(5, 17):
    for j in range(1, len(rcols) + 1):
        c = rl.cell(row=r, column=j)
        c.font = font(italic=(r == 5), color="7F7F7F" if r == 5 else "000000")
        c.border = BOX
        c.alignment = WRAP
        if j != 6:
            c.fill = INPUT_FILL
    rl.cell(row=r, column=4).number_format = "yyyy-mm-dd hh:mm"
    rl.cell(row=r, column=5).number_format = "yyyy-mm-dd hh:mm"
    d = rl.cell(row=r, column=6, value=f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")')
    d.number_format = "[h]:mm"
for j, v in enumerate(ex, 1):
    if v is not None:
        c = rl.cell(row=5, column=j, value=v)
from datetime import datetime  # noqa: E402
rl["D5"] = datetime(2026, 10, 21, 9, 30)
rl["E5"] = datetime(2026, 10, 21, 9, 33)
for col, opts in (("B", "Quick trial,Full replication,Partial (resumed)"),
                  ("G", "Finished,Failed,Interrupted"),
                  ("H", "REPRODUCED EXACTLY,REPRODUCED,NOT REPRODUCED,N/A (trial)")):
    v = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    rl.add_data_validation(v)
    v.add(f"{col}5:{col}16")

st0 = 19
rl.cell(row=st0, column=1, value="Stage run times (full replication)").font = font(bold=True, color="1F3864")
header(rl, st0 + 1, ["Stage", "Description", "", "Published time (2 cores, 8 GB)", "Your time (h:mm:ss)", "", "Notes"])
stages = [(1, "Build the synthetic populations", "0:01:17"), (2, "Estimate: parametric-rate bootstrap (reported)", "1:08:24"),
          (3, "Estimate: fixed-rate bootstrap (comparison)", "1:12:19"), (4, "Part D pre-IRA reference", "1:08:53"),
          (5, "Join estimates to the truth", "0:00:01"), (6, "Manuscript tables and figures", "0:00:02"),
          (7, "Automated tests", "0:01:02"), (8, "Result check", "0:00:01")]
for k, (n, desc, t) in enumerate(stages):
    r = st0 + 2 + k
    rl.cell(row=r, column=1, value=n).font = font()
    rl.cell(row=r, column=2, value=desc).font = font()
    h, m, s = map(int, t.split(":"))
    c = rl.cell(row=r, column=4, value=(h * 3600 + m * 60 + s) / 86400)
    c.number_format = "[h]:mm:ss"
    c.font = font(color="0000FF")
    y = rl.cell(row=r, column=5)
    y.fill = INPUT_FILL
    y.number_format = "[h]:mm:ss"
    y.border = BOX
    rl.cell(row=r, column=7).fill = INPUT_FILL
tot = st0 + 2 + len(stages)
rl.cell(row=tot, column=2, value="Total").font = font(bold=True)
c = rl.cell(row=tot, column=4, value=f"=SUM(D{st0 + 2}:D{tot - 1})")
c.number_format = "[h]:mm:ss"
c.font = font(bold=True)
c = rl.cell(row=tot, column=5, value=f"=SUM(E{st0 + 2}:E{tot - 1})")
c.number_format = "[h]:mm:ss"
c.font = font(bold=True)
rl.cell(row=tot + 1, column=1,
        value="Published times: measured on the author's test re-run, 2026-10-20 (start-here/README.md, step 7). "
              "Enter your times from run_log.txt ('Stage N finished in ...').").font = font(italic=True, color="595959")
rl.freeze_panes = "A5"

# ---------------- Issues log ----------------
il = wb.create_sheet("Issues log")
title(il, "Issues log", "One row per problem found. Use the ID (I-01, I-02, ...) in the Checklist's Issue ID column.")
icols = ["Issue ID", "Date found", "Checklist ID", "Severity", "Description", "File / location", "How found",
         "Status", "Resolution", "Date resolved", "Resolved by"]
header(il, 4, icols, [9, 12, 11, 11, 50, 36, 26, 11, 40, 13, 14])
example = ["I-01 (EXAMPLE)", datetime(2026, 10, 21), "C21", "Minor",
           "Part C errors/config.yaml design block lists 1,500 enrollees per stratum; the estimation uses 330 per tercile (990).",
           "part-c-ipm/errors/config.yaml (design: strata)", "Code review", "Open",
           "Metadata only (written to design.yaml); no result affected. Author to correct the description.", None, None]
for r in range(5, 40):
    for j in range(1, len(icols) + 1):
        c = il.cell(row=r, column=j)
        c.fill = INPUT_FILL
        c.border = BOX
        c.alignment = WRAP
        c.font = font(italic=(r == 5), color="7F7F7F" if r == 5 else "000000")
    il.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    il.cell(row=r, column=10).number_format = "yyyy-mm-dd"
for j, v in enumerate(example, 1):
    if v is not None:
        il.cell(row=5, column=j, value=v)
for col, opts in (("D", "Critical,Major,Minor,Note"), ("H", "Open,Resolved,Accepted")):
    v = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    il.add_data_validation(v)
    v.add(f"{col}5:{col}39")
il.cell(row=41, column=1, value="Severity: Critical = a reported result is wrong or cannot be reproduced. Major = an assumption "
        "or method differs from its description in a way that could change a result. Minor = documentation or labeling "
        "error with no effect on results. Note = suggestion. Status 'Accepted' = left open by agreement, explained in "
        "Resolution. Row 5 is an example (a real discrepancy, also listed as Checklist C21). If you confirm it, delete "
        "'(EXAMPLE)' from its ID so it is counted; otherwise overwrite the row.").font = font(italic=True, color="595959")
il.cell(row=41, column=1).alignment = WRAP
il.merge_cells("A41:K41")
il.row_dimensions[41].height = 45
il.freeze_panes = "A5"

# ---------------- Summary ----------------
sm = wb.create_sheet("Summary", 1)
title(sm, "Summary", "Calculated from the Checklist and Issues log. Nothing to fill in here.")
header(sm, 4, ["Section", "Tasks", "Pass", "Fail", "N/A", "Not started", "Complete (%)"], [44, 9, 9, 9, 9, 12, 13])
CK = f"Checklist!$B${first_data}:$B${last_data}"
CS = f"Checklist!$G${first_data}:$G${last_data}"
r = 5
for sec, name in SECTIONS.items():
    sm.cell(row=r, column=1, value=f"{sec}. {name}").font = font()
    sm.cell(row=r, column=2, value=f'=COUNTIF({CK},"{name}")')
    for j, s in zip((3, 4, 5, 6), ("Pass", "Fail", "N/A", "Not started")):
        sm.cell(row=r, column=j, value=f'=COUNTIFS({CK},"{name}",{CS},"{s}")')
    c = sm.cell(row=r, column=7, value=f"=IF(B{r}=0,0,(C{r}+D{r}+E{r})/B{r})")
    c.number_format = "0%"
    for j in range(1, 8):
        sm.cell(row=r, column=j).border = BOX
        sm.cell(row=r, column=j).font = font()
    r += 1
sm.cell(row=r, column=1, value="All sections").font = font(bold=True)
for j in range(2, 7):
    L = get_column_letter(j)
    sm.cell(row=r, column=j, value=f"=SUM({L}5:{L}{r - 1})").font = font(bold=True)
c = sm.cell(row=r, column=7, value=f"=IF(B{r}=0,0,(C{r}+D{r}+E{r})/B{r})")
c.number_format = "0%"
c.font = font(bold=True)
tot_row = r
IS = "'Issues log'!$D$5:$D$39"
IST = "'Issues log'!$H$5:$H$39"
IID = "'Issues log'!$A$5:$A$39"
NOEX = f'{IID},"<>*EXAMPLE*"'
r += 2
sm.cell(row=r, column=1, value="Issues (rows whose ID contains EXAMPLE are not counted)").font = font(bold=True, color="1F3864")
header(sm, r + 1, ["Severity", "Total", "Open", "Resolved", "Accepted"])
for k, sev in enumerate(("Critical", "Major", "Minor", "Note")):
    rr = r + 2 + k
    sm.cell(row=rr, column=1, value=sev).font = font()
    sm.cell(row=rr, column=2, value=f'=COUNTIFS({IS},"{sev}",{NOEX})').font = font()
    for j, s in zip((3, 4, 5), ("Open", "Resolved", "Accepted")):
        sm.cell(row=rr, column=j, value=f'=COUNTIFS({IS},"{sev}",{IST},"{s}",{NOEX})').font = font()
crit = r + 2
maj = r + 3
rr = r + 7
sm.cell(row=rr, column=1, value="Fail rows without an Issue ID").font = font()
sm.cell(row=rr, column=2, value=f'=COUNTIFS({CS},"Fail",Checklist!$K${first_data}:$K${last_data},"")').font = font()
rr += 2
sm.cell(row=rr, column=1, value="READY TO SIGN?").font = font(bold=True, size=12)
c = sm.cell(row=rr, column=2, value=(
    f'=IF(AND(F{tot_row}=0,C{crit}=0,C{maj}=0,B{rr - 2}=0),"YES: every task checked; no open Critical or Major issues",'
    f'"NOT YET: "&F{tot_row}&" task(s) not started, "&(C{crit}+C{maj})&" open Critical/Major issue(s), "'
    f'&B{rr - 2}&" Fail row(s) without an Issue ID")'))
c.font = font(bold=True, size=12)
sm.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=7)
sm.cell(row=rr + 1, column=1, value="A Fail row does not by itself block sign-off if its issue is Resolved or Accepted.").font = font(italic=True, color="595959")
sm.column_dimensions["A"].width = 44

for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToHeight = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)}: {len(ROWS)} checklist tasks")
